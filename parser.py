"""Поиск игровых каналов на YouTube и Twitch с контактами и выгрузкой в Excel.

Без API-ключей и авторизации, без LLM. Пример запуска:
    python parser.py --min-subs 5000 --max-subs 200000 --limit 100 --minutes 15
Все параметры: python parser.py --help
"""
import argparse
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")

YT_QUERIES = ["прохождение игр", "летсплей", "игровой канал", "стрим игры", "обзор игр",
              "minecraft", "cs2", "dota 2", "gta 5", "инди игры", "разработка игр", "геймдев",
              "unity уроки", "unreal engine", "ретро игры", "хоррор игры"]
GAMING_WORDS = ["игр", "game", "gaming", "летсплей", "let's play", "прохожд", "стрим", "stream",
                "minecraft", "майнкрафт", "cs2", "cs:go", "dota", "дота", "gta", "fortnite",
                "roblox", "роблокс", "unity", "unreal", "геймдев", "gamedev", "инди", "indie",
                "rpg", "шутер", "хоррор", "horror", "speedrun", "спидран", "twitch", "киберспорт"]
TWITCH_NOT_GAMES = {"Just Chatting", "Music", "Art", "IRL", "Talk Shows & Podcasts", "ASMR",
                    "Sports", "Travel & Outdoors", "Food & Drink", "Pools, Hot Tubs, and Beaches",
                    "Makers & Crafting", "Beauty & Body Art", "Science & Technology", "Politics",
                    "Animals, Aquariums, and Zoos", "Fitness & Health", "Special Events",
                    "Slots", "Virtual Casino"}
TWITCH_CLIENT_ID = "kimne78kx3ncx6brgo4mv6wki5h1ko"  # публичный ID веб-версии Twitch


@dataclass
class Channel:
    platform: str
    name: str
    url: str
    subscribers: int
    avg_views: int | None = None
    emails: list[str] = field(default_factory=list)
    telegram: list[str] = field(default_factory=list)
    vk: list[str] = field(default_factory=list)
    other: list[str] = field(default_factory=list)
    description: str = ""

    @property
    def has_contact(self) -> bool:
        return bool(self.emails or self.telegram or self.vk or self.other)


# ---------- общее ----------

def parse_count(text: str) -> int | None:
    """'1.82M subscribers' -> 1820000, '38K' -> 38000, '1,234 views' -> 1234."""
    m = re.search(r"([\d.,]+)\s*([KMB])?", text.replace(" ", " "))
    if not m:
        return None
    num, mult = m.group(1), m.group(2)
    num = num.replace(",", "") if mult is None else num.replace(",", ".")
    try:
        value = float(num)
    except ValueError:
        return None
    return int(value * {"K": 1e3, "M": 1e6, "B": 1e9}.get(mult or "", 1))


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
TG_RE = re.compile(r"(?:https?://)?(?:t\.me|telegram\.me)/[\w+/-]+", re.I)
VK_RE = re.compile(r"(?:https?://)?(?:vk\.com|vk\.ru)/[\w.-]+", re.I)
OTHER_RE = re.compile(r"(?:https?://)?(?:discord\.gg|discord\.com/invite|boosty\.to|hipolink\.me|"
                      r"taplink\.cc|linktr\.ee|instagram\.com|tiktok\.com)/[\w.@/-]+", re.I)


def extract_contacts(ch: Channel, *texts: str) -> None:
    blob = "\n".join(texts)
    uniq = lambda xs: list(dict.fromkeys(x.rstrip(".,)/") for x in xs))
    link = lambda xs: uniq("https://" + re.sub(r"^https?://", "", x, flags=re.I) for x in xs)
    ch.emails = uniq(e.lower() for e in EMAIL_RE.findall(blob) if not e.endswith((".png", ".jpg")))
    ch.telegram = link(TG_RE.findall(blob))
    ch.vk = link(VK_RE.findall(blob))
    ch.other = link(OTHER_RE.findall(blob))


def is_gaming(*texts: str) -> bool:
    blob = " ".join(texts).lower()
    return any(w in blob for w in GAMING_WORDS)


def count(found: dict, platform: str) -> int:
    return sum(1 for key in found if key[0] == platform)


def walk(obj, key, out):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k == key:
                out.append(v)
            walk(v, key, out)
    elif isinstance(obj, list):
        for v in obj:
            walk(v, key, out)
    return out


# ---------- YouTube ----------

class YouTube:
    def __init__(self, pause: float):
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9"})
        self.s.cookies.set("SOCS", "CAI", domain=".youtube.com")  # пропустить окно согласия cookie
        self.pause = pause
        self.key = self.version = None

    def _initial_data(self, html: str) -> dict:
        if self.key is None:
            self.key = re.search(r'"INNERTUBE_API_KEY":"([^"]+)"', html).group(1)
            self.version = re.search(r'"clientVersion":"([^"]+)"', html).group(1)
        return json.loads(re.search(r"var ytInitialData = (\{.*?\});</script>", html).group(1))

    def search_channels(self, query: str, pages: int):
        """Выдаёт channelRenderer из поиска с фильтром «Каналы»."""
        html = self.s.get("https://www.youtube.com/results",
                          params={"search_query": query, "sp": "EgIQAg==", "hl": "en"}, timeout=30).text
        data = self._initial_data(html)
        for page in range(pages):
            yield from walk(data, "channelRenderer", [])
            tokens = [c.get("token") for c in walk(data, "continuationCommand", []) if c.get("token")]
            if not tokens or page == pages - 1:
                return
            time.sleep(self.pause)
            data = self.s.post(f"https://www.youtube.com/youtubei/v1/search?key={self.key}", timeout=30, json={
                "context": {"client": {"clientName": "WEB", "clientVersion": self.version, "hl": "en"}},
                "continuation": tokens[-1]}).json()

    def details(self, ch: Channel, channel_id: str, with_views: bool) -> None:
        html = self.s.get(f"https://www.youtube.com/channel/{channel_id}/videos", params={"hl": "en"},
                          timeout=30).text
        data = self._initial_data(html)
        ch.description = data.get("metadata", {}).get("channelMetadataRenderer", {}).get("description", "")
        links = [requests.utils.unquote(u) for u in re.findall(r"q=(https?%3A%2F%2F[^\"&\\]+)", html)]
        extract_contacts(ch, ch.description, *links)
        if with_views:
            views = [parse_count(v) for v in re.findall(
                r'"content":"([\d.,]+[KMB]?)"\},"accessibilityLabel":"[^"]*views"', html)][:10]
            views = [v for v in views if v is not None]
            ch.avg_views = int(sum(views) / len(views)) if views else None


def run_youtube(args, found: dict, deadline: float) -> None:
    yt = YouTube(args.pause)
    for query in args.queries:
        for r in yt.search_channels(query, args.pages):
            if time.time() > deadline or count(found, "yt") >= args.limit:
                return
            cid = r.get("channelId")
            if not cid or ("yt", cid) in found:
                continue
            # в выдаче YouTube число подписчиков лежит в videoCountText, а в subscriberCountText — @handle
            subs = parse_count(r.get("videoCountText", {}).get("simpleText", "") or "")
            if subs is None or not (args.min_subs <= subs <= args.max_subs):
                continue
            title = r.get("title", {}).get("simpleText", "")
            handle = r.get("subscriberCountText", {}).get("simpleText", "")
            url = f"https://www.youtube.com/{handle}" if handle.startswith("@") else \
                f"https://www.youtube.com/channel/{cid}"
            ch = Channel("YouTube", title, url, subs)
            time.sleep(args.pause)
            try:
                yt.details(ch, cid, not args.no_views)
            except Exception as e:
                print(f"  ! {title}: {e}", file=sys.stderr)
                continue
            if not args.any_topic and not is_gaming(title, ch.description, query):
                continue
            if args.require_contact and not ch.has_contact:
                continue
            found[("yt", cid)] = ch
            print(f"[YouTube] {len(found):>4}. {title} | {subs:,} подп. | {', '.join(ch.emails + ch.telegram)[:60]}")


# ---------- Twitch ----------

class Twitch:
    def __init__(self, pause: float):
        self.s = requests.Session()
        self.s.headers.update({"Client-ID": TWITCH_CLIENT_ID, "User-Agent": UA})
        self.pause = pause

    def gql(self, query: str) -> dict:
        r = self.s.post("https://gql.twitch.tv/gql", data=json.dumps({"query": query}), timeout=30)
        return r.json().get("data") or {}

    def top_games(self, n: int) -> list[str]:
        data = self.gql("query { games(first: %d) { edges { node { name } } } }" % min(n + 20, 100))
        names = [e["node"]["name"] for e in data.get("games", {}).get("edges", [])]
        return [g for g in names if g not in TWITCH_NOT_GAMES][:n]

    def streams(self, game: str, languages: list[str], pages: int):
        cursor = None
        langs = ",".join(f'"{l.upper()}"' for l in languages)
        for _ in range(pages):
            after = f', after: "{cursor}"' if cursor else ""
            data = self.gql('query { game(name: %s) { streams(first: 30%s, options: {languages: [%s]}) '
                            '{ edges { cursor node { viewersCount broadcaster { id login displayName '
                            'description followers { totalCount } } } } } } }'
                            % (json.dumps(game), after, langs))
            edges = ((data.get("game") or {}).get("streams") or {}).get("edges") or []
            yield from (e["node"] for e in edges)
            if len(edges) < 30:
                return
            cursor = edges[-1]["cursor"]
            time.sleep(self.pause)


def run_twitch(args, found: dict, deadline: float) -> None:
    tw = Twitch(args.pause)
    games = args.games or tw.top_games(args.twitch_games)
    for game in games:
        for node in tw.streams(game, args.lang, args.pages):
            if time.time() > deadline or count(found, "tw") >= args.limit:
                return
            b = node.get("broadcaster") or {}
            if not b or ("tw", b["id"]) in found:
                continue
            subs = (b.get("followers") or {}).get("totalCount") or 0
            if not (args.min_subs <= subs <= args.max_subs):
                continue
            ch = Channel("Twitch", b["displayName"], f"https://www.twitch.tv/{b['login']}", subs,
                         avg_views=node.get("viewersCount"), description=b.get("description") or "")
            extract_contacts(ch, ch.description)
            if args.require_contact and not ch.has_contact:
                continue
            found[("tw", b["id"])] = ch
            print(f"[Twitch]  {len(found):>4}. {ch.name} ({game}) | {subs:,} фолл. | {', '.join(ch.emails + ch.telegram)[:60]}")


# ---------- Excel ----------

def save_xlsx(channels: list[Channel], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Каналы"
    head = ["Платформа", "Канал", "Ссылка", "Подписчики / фолловеры", "Средние просмотры",
            "Почта", "Telegram", "VK", "Другие контакты", "Описание"]
    ws.append(head)
    for c in channels:
        ws.append([c.platform, c.name, c.url, c.subscribers, c.avg_views, "\n".join(c.emails),
                   "\n".join(c.telegram), "\n".join(c.vk), "\n".join(c.other), c.description[:500]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate([10, 28, 42, 14, 14, 30, 30, 30, 30, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description="Поиск игровых каналов YouTube и Twitch с контактами")
    p.add_argument("--platforms", default="youtube,twitch", help="youtube,twitch или одно из них")
    p.add_argument("--min-subs", type=int, default=1000, help="минимум подписчиков (по умолчанию 1000)")
    p.add_argument("--max-subs", type=int, default=500_000, help="максимум подписчиков")
    p.add_argument("--limit", type=int, default=100, help="сколько каналов найти на каждой платформе")
    p.add_argument("--minutes", type=float, default=20, help="лимит времени работы в минутах")
    p.add_argument("--queries", nargs="+", default=YT_QUERIES, help="поисковые запросы YouTube")
    p.add_argument("--games", nargs="+", help="категории Twitch (по умолчанию топ игровых)")
    p.add_argument("--twitch-games", type=int, default=25, help="сколько топ-категорий Twitch взять")
    p.add_argument("--lang", nargs="+", default=["ru"], help="язык стримов Twitch: ru en ...")
    p.add_argument("--pages", type=int, default=5, help="страниц выдачи на запрос / категорию")
    p.add_argument("--no-views", action="store_true", help="не считать средние просмотры YouTube (быстрее)")
    p.add_argument("--any-topic", action="store_true", help="не проверять игровую тематику YouTube")
    p.add_argument("--require-contact", action=argparse.BooleanOptionalAction, default=True,
                   help="брать только каналы с контактом (по умолчанию да)")
    p.add_argument("--pause", type=float, default=0.7, help="пауза между запросами, сек")
    p.add_argument("--out", default="channels.xlsx", help="имя файла Excel")
    p.add_argument("--no-open", action="store_true", help="не открывать Excel после завершения")
    args = p.parse_args()
    if len(sys.argv) == 1:
        ask(args)

    start = time.time()
    found: dict = {}
    platforms = [x.strip().lower() for x in args.platforms.split(",")]
    try:
        if "twitch" in platforms:
            # время делится между платформами, чтобы второй тоже хватило
            share = 0.5 if "youtube" in platforms else 1.0
            run_twitch(args, found, start + args.minutes * 60 * share)
        if "youtube" in platforms:
            run_youtube(args, found, start + args.minutes * 60)
    except KeyboardInterrupt:
        print("\nОстановлено, сохраняю то, что нашёл")
    channels = sorted(found.values(), key=lambda c: (c.platform, -c.subscribers))
    save_xlsx(channels, args.out)
    print(f"\nГотово: {len(channels)} каналов -> {args.out}")
    if not args.no_open and hasattr(os, "startfile"):
        os.startfile(os.path.abspath(args.out))


def ask(args) -> None:
    """Запуск без параметров: спрашиваем основное, Enter оставляет значение по умолчанию."""
    def q(prompt, default, cast=str):
        raw = input(f"{prompt} [{default}]: ").strip()
        return cast(raw) if raw else default
    print("Поиск игровых каналов. Enter = значение по умолчанию.\n")
    args.platforms = q("Платформы (youtube,twitch)", args.platforms)
    args.min_subs = q("Минимум подписчиков", args.min_subs, int)
    args.max_subs = q("Максимум подписчиков", args.max_subs, int)
    args.limit = q("Сколько каналов найти на каждой платформе", args.limit, int)
    args.minutes = q("Лимит времени, минут", args.minutes, float)
    args.out = q("Имя файла Excel", args.out)
    print()


if __name__ == "__main__":
    main()
